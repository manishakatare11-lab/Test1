import os.path
import time

import openpyxl
from selenium.common import NoSuchElementException
from selenium.webdriver.common.by import By
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from utilities.jsonutil import read_json
from openpyxl import Workbook, load_workbook



class BasePage:

    def __init__(self, driver, json_path):
        self.driver = driver
        self.locators = read_json(json_path)

    def get_by(self, locator_type):
        locator_type = locator_type.lower()

        if locator_type == "id":
            return By.ID
        elif locator_type == "name":
            return By.NAME
        elif locator_type == "xpath":
            return By.XPATH
        elif locator_type == "css":
            return By.CSS_SELECTOR
        elif locator_type == "class":
            return By.CLASS_NAME
        elif locator_type == "linktext":
            return By.LINK_TEXT
        elif locator_type == "partiallinktext":
            return By.PARTIAL_LINK_TEXT
        else:
            raise Exception(f"Locator type not supported: {locator_type}")

    def get_element(self, element_name, timeout=20):
        element = self.locators.get(element_name)

        if element is None:
            raise Exception(f"Element '{element_name}' not found in locators")

        by = self.get_by(element["find_by"])
        value = element["identifier"]

        return WebDriverWait(self.driver, timeout).until(
            EC.visibility_of_element_located((by, value))
        )

    def do_send_keys(self, element_name, text):
        element = self.locators.get(element_name)

        if element is None:
            raise Exception(f"Element '{element_name}' not found in locators")

        by = self.get_by(element["find_by"])
        locator = element["identifier"]

        element = WebDriverWait(self.driver, 20).until(
            EC.visibility_of_element_located((by, locator))
        )

        element.clear()
        element.click()
        element.send_keys(text)

    def do_click(self, element_name):
        element = self.locators.get(element_name)

        by = self.get_by(element["find_by"])
        locator = element["identifier"]

        WebDriverWait(self.driver, 20).until(
            EC.element_to_be_clickable((by, locator))
        ).click()

    def get_text(self, element_name):
        return self.get_element(element_name).text

    def wait_for_page_load(self, timeout=30):
        WebDriverWait(self.driver, timeout).until(
            lambda driver: driver.execute_script("return document.readyState") == "complete"
        )

    def do_select_DD(self, dropdown_name, index):

        element = self.locators.get(dropdown_name)

        if element is None:
            raise Exception(f"Element '{dropdown_name}' not found in locators")

        by = self.get_by(element["find_by"])
        locator = element["identifier"]

        # Open dropdown
        WebDriverWait(self.driver, 20).until(
            EC.element_to_be_clickable((by, locator))
        ).click()

        # Select option

        option_xpath = f"(//a[@class='option'])[{index}]"

        WebDriverWait(self.driver, 20).until(
            EC.element_to_be_clickable((By.XPATH, option_xpath))
        ).click()

    def scroll_down_until_element(self, locator, max_scroll=10):

        for i in range(max_scroll):

            try:
                element = self.driver.find_element(*locator)

                if element.is_displayed():
                    return element

            except:
                pass

            self.driver.execute_script("window.scrollBy(0,500);")
            time.sleep(1)

    def is_element_visible(self, element_name, timeout=5):
        """
        Check whether element is visible or not
        """

        try:
            element = self.locators.get(element_name)

            if element is None:
                raise Exception(f"Element '{element_name}' not found in locators")

            by = self.get_by(element["find_by"])
            locator = element["identifier"]

            WebDriverWait(self.driver, timeout).until(
                EC.visibility_of_element_located((by, locator))
            )

            return True

        except:
            return False

    def Save_details(self, candidate_name, candidate_email,
                     common_login_url,
                     file_path="../utilities/details.xlsx"):

        # Check file exists or not
        if os.path.exists(file_path):
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active

        else:
            workbook = Workbook()
            sheet = workbook.active

            # Header row
            sheet.append([
                "Candidate Name",
                "Candidate Email",
                "Common Login URL"
            ])

        # Add data row
        sheet.append([
            candidate_name,
            candidate_email,
            common_login_url
        ])

        workbook.save(file_path)

        print(f"{candidate_name} | {candidate_email} | {common_login_url} saved")

    def read_excel(self):
        workbook = openpyxl.load_workbook("../utilities/details.xlsx")
        sheet = workbook.active

        data =[]
        for row in range(2, sheet.max_row + 1):
            candidate_name = sheet.cell(row=row, column=1).value
            candidate_email = sheet.cell(row=row, column=2).value
            common_login_url = sheet.cell(row=row, column=3).value
            data.append((candidate_name, candidate_email, common_login_url))

        return data

    def accept_cookies(self):
        try:
            wait = WebDriverWait(self.driver, 10)
            ok_button = wait.until(
                EC.element_to_be_clickable((By.ID, "CybotCookiebotDialogBodyLevelButtonAccept"))
            )
            ok_button.click()
        except:
            pass